"""Inputs that are not in the sample's exact form: scanned or photographed artwork (read by OCR),
and content sheets in other layouts (CSV, header-row tables, other label names, UK nutrient names).

The scan is the Esko carton flattened to a 300 dpi picture of its print view (no live text, no
dieline guides), as a scanner would see the printed carton.
"""

from pathlib import Path

import openpyxl
import pymupdf
import pytest
from fastapi.testclient import TestClient

import app as api
import runs
from checker.artwork import ArtworkError, read_artwork
from checker.artwork.ocr import fix_digits
from checker.artwork.render import _print_view
from checker.content_sheet import SheetError, read_content_sheet
from checker.feedback import FeedbackStore
from checker.models import Verdict
from checker.nutrition import read_nutrition_panel
from checker.pdf_work import InlinePdfWork
from checker.pipeline import OCR_CAVEAT, OCR_INDEPENDENT_CHECKS, ReviewPipeline
from checker.sections import find_blocks

SAMPLES = Path(__file__).resolve().parents[2] / "samples"
RULES = Path(__file__).resolve().parents[2] / "rules"
CARTON = SAMPLES / "DancingFish_Granola_Berry_US_Carton.pdf"
SHEET = SAMPLES / "DancingFish_Granola_BerryCrunch_US_Content_Sheet.xlsx"


@pytest.fixture(scope="module")
def scan_png(tmp_path_factory) -> Path:
    path = tmp_path_factory.mktemp("scan") / "carton_scan.png"
    path.write_bytes(_print_view(CARTON, 300, 0))
    return path


@pytest.fixture(scope="module")
def scan_pdf(scan_png) -> Path:
    rect = pymupdf.open(CARTON)[0].rect
    doc = pymupdf.open()
    page = doc.new_page(width=rect.width, height=rect.height)
    page.insert_image(page.rect, filename=str(scan_png))
    path = scan_png.with_name("carton_scan.pdf")
    doc.save(path)
    return path


@pytest.fixture(scope="module")
def scan_review(scan_pdf, tmp_path_factory):
    store = FeedbackStore(tmp_path_factory.mktemp("fb") / "feedback.json")
    return ReviewPipeline(feedback=store).run(scan_pdf, SHEET, [], use_ai=False)


# ---------------- OCR ----------------


def test_a_scan_is_read_by_ocr(scan_pdf):
    art = read_artwork(scan_pdf)
    assert art.text_source == "ocr" and len(art.lines) > 80
    nf = read_nutrition_panel(art)
    assert nf.calories == 166
    amounts = {k: v.amount for k, v in nf.nutrients.items()}
    assert (amounts["total_fat"], amounts["total_carbohydrate"], amounts["protein"]) == (15, 66, 34)
    ingredients = find_blocks(art)["ingredients"].text
    assert "Wholegrain Oat Flakes" in ingredients and "Honey" in ingredients
    assert art.barcodes and art.barcodes[0]["text"].endswith("20000000036")  # barcodes come from the image anyway


def test_problems_in_ocr_text_are_likely_never_definite(scan_review):
    findings = scan_review.report.findings
    calories = next(f for f in findings if f.check == "calories_atwater")
    assert calories.verdict == Verdict.LIKELY and OCR_CAVEAT in calories.explanation
    assert not [f.title for f in findings if f.verdict == Verdict.DEFINITE and f.check not in OCR_INDEPENDENT_CHECKS]
    assert "read by OCR" in scan_review.ai_note and "one panel" in scan_review.ai_note


def test_missing_items_on_a_scan_are_not_proven_missing(scan_pdf, tmp_path):
    """"No FSSAI licence number" cites no line, so a line-based cap missed it: OCR may simply not have read it."""
    rules = [RULES / "india_fssai_food.json", RULES / "artwork_approval_core.json"]
    result = ReviewPipeline(feedback=FeedbackStore(tmp_path / "fb.json")).run(scan_pdf, SHEET, rules, use_ai=False)
    licence = next(f for f in result.report.findings if f.check == "fssai_licence")
    assert licence.verdict == Verdict.LIKELY and OCR_CAVEAT in licence.explanation
    assert not [f.title for f in result.report.findings if f.verdict == Verdict.DEFINITE and f.check not in OCR_INDEPENDENT_CHECKS]


def test_a_searchable_scan_is_treated_as_ocr_text(tmp_path):
    """Scanning software lays its own OCR text over the picture, invisibly. That text is a guess,
    so it must not prove anything the way real artwork text does."""
    doc = pymupdf.open()
    page = doc.new_page(width=400, height=300)
    page.insert_image(page.rect, pixmap=pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 30), 0))
    for y, text in [(60, "Nutrition Facts"), (80, "Serving size 30g"), (100, "Calories 900"),
                    (120, "Total Fat 1g"), (140, "Total Carbohydrate 2g"), (160, "Protein 1g")]:
        page.insert_text((40, y), text, fontsize=10, render_mode=3)
    path = tmp_path / "searchable.pdf"
    doc.save(path)
    assert read_artwork(path).text_source == "ocr_layer"
    result = ReviewPipeline(feedback=FeedbackStore(tmp_path / "fb.json")).run(path, None, [], use_ai=False)
    calories = next(f for f in result.report.findings if f.check == "calories_atwater")
    assert calories.verdict == Verdict.LIKELY and "invisible OCR layer" in result.ai_note


def test_an_ocr_failure_is_an_error_not_an_empty_review(tmp_path, monkeypatch):
    import checker.artwork.reader as reader

    def broken(page):
        raise RuntimeError("onnxruntime failed")
    monkeypatch.setattr(reader, "ocr_lines", broken)
    doc = pymupdf.open()
    doc.new_page(width=200, height=120).draw_rect(pymupdf.Rect(10, 10, 190, 110), fill=(0.2, 0.4, 0.9))
    path = tmp_path / "unreadable.pdf"
    doc.save(path)
    with pytest.raises(ArtworkError, match="OCR could not read unreadable.pdf"):
        read_artwork(path)


def test_oversized_and_bomb_images_are_refused_before_decoding():
    import io

    from PIL import Image

    from checker.artwork import image_to_pdf
    bomb, big = io.BytesIO(), io.BytesIO()
    Image.new("1", (20000, 20000), 1).save(bomb, format="PNG")  # 88 KB of file, 400 megapixels
    Image.new("1", (8000, 7500), 1).save(big, format="PNG")  # 60 megapixels: allowed by Pillow, too big here
    with pytest.raises(ArtworkError, match="too large to open safely"):
        image_to_pdf(bomb.getvalue(), "bomb.png")  # before the fix: an unhandled DecompressionBombError (a 500)
    with pytest.raises(ArtworkError, match="8000 x 7500 pixels.*50 megapixels"):
        image_to_pdf(big.getvalue(), "big.png")


def test_the_scan_finds_what_the_pdf_finds(scan_review):
    """Same code checks, same verdicts, except where the dieline (panels, glyph sizes) is needed."""
    live = ReviewPipeline(feedback=FeedbackStore(Path(scan_review.report.artwork).parent / "x.json")).run(CARTON, SHEET, [], use_ai=False)
    needs_dieline = {"net_quantity_pdp", "net_qty_height", "type_size", "no_added_sugar"}
    live_issues = {f.check for f in live.report.findings if f.verdict.value in {"definite", "likely"}} - needs_dieline
    scan_issues = {f.check for f in scan_review.report.findings if f.verdict.value in {"definite", "likely"}}
    assert live_issues <= scan_issues, live_issues - scan_issues


def test_label_photo_amounts_are_not_mistaken_for_the_net_weight():
    """Seen on real Cheerios and Oreo label photos: OCR split the title into 'Nutrition' / 'Facts' and read
    'Total Fat' and its '1.5g' as separate pieces far apart, so '1.5g' was reported as the net weight."""
    from checker.models import Artwork, BBox, TextLine
    from checker.nutrition import nutrition_title
    from checker.sections import net_quantity_lines

    def line(i, text, x, y, size=9.0):
        return TextLine(id=f"L{i:03d}", text=text, bbox=BBox(x0=x, y0=y, x1=x + 6 * len(text), y1=y + size),
                        size=size, font="OCR", panel="front")
    lines = [line(0, "Nutrition", 10, 10, 20), line(1, "Facts", 10, 32, 20),
             line(2, "Total Fat", 10, 80), line(3, "1.5g", 200, 80),  # name and amount far apart on one row
             line(4, "Sat Fat", 10, 300), line(5, "Less than", 60, 300), line(6, "20g", 120, 300),  # footnote
             line(7, "NET WT 12 OZ (340g)", 400, 500)]
    art = Artwork(file_name="photo.pdf", width=600, height=600, panels={"front": BBox(x0=0, y0=0, x1=600, y1=600)},
                  lines=lines, text_source="ocr")
    assert nutrition_title(art).text == "Nutrition"
    assert [ln.text for ln in net_quantity_lines(art)] == ["NET WT 12 OZ (340g)"]


def test_ocr_digit_fixes_are_only_the_safe_ones():
    assert fix_digits("Incl. Og Added Sugars") == "Incl. 0g Added Sugars"
    assert fix_digits("Calories 1O0") == "Calories 100"
    assert fix_digits("Sodium l5mg") == "Sodium 15mg"
    assert fix_digits("%8") == "8%"
    assert fix_digits("Organic Oats") == "Organic Oats"  # words are never touched
    assert fix_digits("GOOD FOOD") == "GOOD FOOD"


def test_a_blank_scan_says_so(tmp_path):
    doc = pymupdf.open()
    doc.new_page(width=300, height=300).draw_rect(pymupdf.Rect(20, 20, 280, 280), fill=(0.9, 0.9, 1))
    path = tmp_path / "blank.pdf"
    doc.save(path)
    with pytest.raises(ArtworkError, match="OCR found no readable text"):
        read_artwork(path)


def test_a_photo_upload_is_reviewed(scan_png, tmp_path, monkeypatch):
    store = FeedbackStore(tmp_path / "feedback.json")
    monkeypatch.setattr(runs, "RUNS_DIR", tmp_path / "runs")
    monkeypatch.setattr(api, "feedback_store", store)
    monkeypatch.setattr(api, "pdf_work", InlinePdfWork())
    monkeypatch.setattr(api, "pipeline", ReviewPipeline(feedback=store, pdf=api.pdf_work))
    client = TestClient(api.app)
    response = client.post("/api/review", data={"use_ai": "false"}, files={
        "artwork": ("carton photo.png", scan_png.read_bytes(), "image/png"),
        "content_sheet": (SHEET.name, SHEET.read_bytes(), "application/octet-stream"),
    })
    assert response.status_code == 200, response.text
    run = response.json()
    assert run["result"]["report"]["artwork"] == "carton_photo.pdf"
    assert "read by OCR" in run["result"]["ai_note"]
    assert client.get(run["image_url"]).status_code == 200


# ---------------- content sheets ----------------


def test_the_esko_sheet_reads_as_before():
    spec = read_content_sheet(SHEET)
    assert len(spec.fields) == 36 and spec.missing_fields() == []
    assert spec.get("ingredient_list").cell == "B26" and spec.get("upc_barcode").cell == "E12"
    assert spec.nutrition_basis == "250g" and len(spec.nutrition) == 15
    assert (spec.nutrition[0].name, spec.nutrition[0].amount, spec.nutrition[0].cell) == ("Calories", "166", "B47")


def test_a_two_column_csv_with_other_label_names(tmp_path):
    path = tmp_path / "spec.csv"
    path.write_text(
        "Field;Value\n"
        "Product;Berry Granola\n"
        "Legal name;Granola with berries\n"
        "Ingredients;Oats, Honey (9%), Almonds\n"
        "Allergens;Contains almonds\n"
        "Net Wt;480 g\n"
        "EAN;0120000000036\n"
        "\n"
        "Nutrient;Per 100 g;%RI\n"
        "Energy;1868kJ/445kcal;22%\n"
        "Fat;15 g;21%\n"
        "of which saturates;5 g;25%\n"
        "Carbohydrate;66 g;25%\n"
        "Fibre;3 g;\n"
        "Protein;34 g;68%\n", encoding="utf-8")
    spec = read_content_sheet(path)
    assert spec.value("ingredient_list") == "Oats, Honey (9%), Almonds"
    assert spec.value("allergen_information") == "Contains almonds"
    assert spec.value("package_size_net_weight") == "480 g" and spec.get("package_size_net_weight").cell == "B6"
    assert spec.value("upc_barcode") == "0120000000036"
    assert spec.value("legal_name_statement_of_identity") == "Granola with berries"
    rows = {n.name: (n.amount, n.dv) for n in spec.nutrition}
    assert rows["Calories"] == ("445", "22%") and rows["Total Fat"] == ("15 g", "21%")
    assert rows["Saturated Fat"][0] == "5 g" and rows["Dietary Fiber"] == ("3 g", None)
    assert spec.missing_fields() == []


def test_a_header_row_table_and_a_separate_nutrition_sheet(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product"
    ws.append(["Product Name", "Ingredient Statement", "Net Weight (g)", "GTIN"])
    ws.append(["Berry Granola", "Oats, Honey, Almonds", "480g", "0120000000036"])
    nut = wb.create_sheet("Nutrition")
    nut.append(["Nutrient", "Amount", "Unit", "% DV"])
    for name, amount, unit, dv in [("Calories", 166, "kcal", None), ("Total Fat", 15, "g", 0.19), ("Sodium", 23, "mg", 0.01)]:
        nut.append([name, amount, unit, dv])
        if dv is not None:
            nut.cell(nut.max_row, 4).number_format = "0%"
    path = tmp_path / "spec.xlsx"
    wb.save(path)
    spec = read_content_sheet(path)
    assert spec.value("product_name") == "Berry Granola" and spec.get("product_name").cell == "A2"
    assert spec.value("ingredient_list") == "Oats, Honey, Almonds"
    assert spec.value("package_size_net_weight") == "480g" and spec.value("upc_barcode") == "0120000000036"
    rows = {n.name: n for n in spec.nutrition}
    assert rows["Total Fat"].amount == "15g" and rows["Total Fat"].dv == "19%" and rows["Total Fat"].cell == "Nutrition!B3"
    assert rows["Calories"].amount == "166"  # calories are a bare number, as on the Esko sheet
    assert spec.missing_fields() == ["statement of identity (legal name)", "allergen information"]


def test_label_and_value_in_one_cell(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B2"] = "INGREDIENTS: Oats, Honey"
    ws["B4"] = "Allergy advice: Contains almonds"
    ws["B6"] = "Net quantity: 480g"
    path = tmp_path / "notes.xlsx"
    wb.save(path)
    spec = read_content_sheet(path)
    assert spec.value("ingredient_list") == "Oats, Honey" and spec.get("ingredient_list").cell == "B2"
    assert spec.value("allergen_information") == "Contains almonds"
    assert spec.value("package_size_net_weight") == "480g"


def test_missing_sheet_fields_are_reported_not_passed(tmp_path):
    path = tmp_path / "thin.csv"
    path.write_text("Product name,Berry Granola\nBrand,Dancing Fish\n", encoding="utf-8")
    result = ReviewPipeline(feedback=FeedbackStore(tmp_path / "fb.json")).run(CARTON, path, [], use_ai=False)
    assert "found no statement of identity (legal name), ingredient list, allergen information, net weight, nutrition table" in result.ai_note


def test_old_and_unreadable_sheets_say_what_to_do(tmp_path):
    old = tmp_path / "spec.xls"
    old.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(SheetError, match="save it as .xlsx"):
        read_content_sheet(old)
    empty = tmp_path / "empty.csv"
    empty.write_text("just some words\n", encoding="utf-8")
    with pytest.raises(SheetError, match="no labelled product fields"):
        read_content_sheet(empty)
