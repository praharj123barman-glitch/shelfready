"""Read the approved product content sheet into labelled fields.

Content sheets come in many layouts, so the reader does not assume one:

* form layouts: a label with its value to the right and an optional note after
  it, in any columns (Esko's sheet has two label columns, A and D);
* tables with a header row ("Product Name | Ingredients | Net Weight") and the
  values in the row below;
* "Label: value" written in one cell;
* the nutrition table under a "Nutrient Composition" heading, or else any rows
  that start with a nutrient name ("Energy", "Fat", "Fibre"...);
* .xlsx / .xlsm (every worksheet) and .csv.

Labels are matched to the names the checks use (``FIELD_ALIASES``), so
"Ingredients", "Ingredient Statement" and "INGREDIENTS:" all become
``ingredient_list``. Every value keeps its cell, so findings point at the exact
approved cell, and ``missing_fields()`` says which important fields were not
found, so a skipped comparison is reported instead of passing silently.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field

SHEET_SUFFIXES = {".xlsx", ".xlsm", ".csv"}
NUTRITION_HEADING = "nutrient composition"

# the field names the checks read, with the other labels sheets use for them (compared after field_key())
FIELD_ALIASES: dict[str, list[str]] = {
    "legal_name_statement_of_identity": [
        "statement_of_identity", "legal_name", "legal_denomination", "name_of_the_food", "name_of_food",
        "regulated_product_name", "sales_name", "legal_product_name", "common_name", "generic_name",
        "product_description", "description_of_food"],
    "product_name": ["product", "item_name", "sku_name", "product_title", "trade_name"],
    "brand": ["brand_name"],
    "package_size_net_weight": [
        "net_weight", "net_wt", "net_quantity", "net_contents", "net_content", "net_qty", "pack_size",
        "package_size", "net_weight_g", "net_mass", "quantity", "weight"],
    "upc_barcode": ["upc", "barcode", "ean", "ean_13", "ean13", "gtin", "gtin_13", "gtin_12", "upc_a",
                    "upc_code", "barcode_number", "ean_code"],
    "ingredient_list": ["ingredients", "ingredient_statement", "ingredients_list", "ingredient_declaration",
                        "list_of_ingredients", "ingredient"],
    "storage_instructions": ["storage", "storage_conditions", "storage_instruction", "how_to_store",
                             "storage_advice", "keep"],
    "marketing_copy": ["tagline", "slogan", "marketing_text", "marketing_message", "front_of_pack_copy",
                       "brand_copy", "marketing"],
    "allergen_information": ["allergens", "allergen_statement", "allergen_info", "allergy_advice",
                             "allergen_declaration", "allergen_advice", "allergy_information", "contains"],
    "dietary_information": ["dietary", "dietary_statement", "dietary_claims", "suitable_for", "diet",
                            "dietary_info"],
    "front_back_panel_claims_certifications": [
        "claims", "claims_certifications", "claims_and_certifications", "certifications", "on_pack_claims",
        "product_claims", "logos_certifications", "logos_and_claims", "claims_logos"],
    "country_of_market": ["market", "country_of_sale", "target_market", "sales_country", "country",
                          "market_country", "country_of_sale_market"],
    "product_type": ["category", "product_category"],
    "serving_size": ["serving", "portion_size", "serving_size_g"],
    "number_of_servings_per_package": ["servings_per_container", "servings_per_package", "servings",
                                       "number_of_servings"],
}
# reported when absent: the comparisons that rest on them are skipped
IMPORTANT_FIELDS = {
    "legal_name_statement_of_identity": "statement of identity (legal name)",
    "ingredient_list": "ingredient list",
    "allergen_information": "allergen information",
    "package_size_net_weight": "net weight",
}

# nutrient labels -> the names the nutrition checks know
NUTRIENT_ALIASES: dict[str, list[str]] = {
    "Calories": ["calories", "energy", "energy_kcal", "calories_kcal", "kcal", "calorie", "energy_value"],
    "Total Fat": ["total_fat", "fat", "fats", "total_fats", "fat_total"],
    "Saturated Fat": ["saturated_fat", "saturates", "of_which_saturates", "saturated", "sat_fat",
                      "saturated_fats", "saturated_fatty_acids"],
    "Trans Fat": ["trans_fat", "trans_fats", "trans", "trans_fatty_acids"],
    "Cholesterol": ["cholesterol"],
    "Sodium": ["sodium"],
    "Total Carbohydrate": ["total_carbohydrate", "carbohydrate", "carbohydrates", "total_carbohydrates",
                           "carbs", "total_carbs", "carbohydrate_total"],
    "Dietary Fiber": ["dietary_fiber", "fiber", "fibre", "dietary_fibre", "total_dietary_fiber"],
    "Total Sugars": ["total_sugars", "sugars", "of_which_sugars", "sugar", "total_sugar"],
    "Incl. Added Sugars": ["incl_added_sugars", "added_sugars", "includes_added_sugars", "added_sugar",
                           "incl_added_sugar"],
    "Protein": ["protein", "proteins"],
    "Vitamin D": ["vitamin_d", "vit_d"],
    "Calcium": ["calcium"],
    "Iron": ["iron"],
    "Potassium": ["potassium"],
}
AMOUNT = re.compile(r"^[<~]?\s*\d+(?:[.,]\d+)?\s*(?:g|mg|mcg|µg|ug|kcal|kj|cal|iu|%)?\s*(?:/.*)?$", re.I)
UNIT = re.compile(r"^(g|mg|mcg|µg|ug|kcal|kj|iu)$", re.I)


class SheetError(ValueError):
    """The content sheet is missing or cannot be read."""


class SheetField(BaseModel):
    label: str  # as written on the sheet, e.g. "Ingredients"
    value: str
    cell: str  # "B26", or "Sheet2!B4" on a later worksheet
    note: str | None = None


class SheetNutrient(BaseModel):
    name: str
    amount: str
    dv: str | None
    cell: str


class ApprovedSpec(BaseModel):
    file_name: str
    sheet_name: str
    fields: dict[str, SheetField]
    nutrition_basis: str | None = None
    nutrition: list[SheetNutrient] = Field(default_factory=list)

    def get(self, key: str) -> SheetField | None:
        return self.fields.get(key)

    def value(self, key: str) -> str | None:
        f = self.fields.get(key)
        return f.value if f else None

    def missing_fields(self) -> list[str]:
        """Important approved values this sheet does not have (by their plain names)."""
        missing = [name for key, name in IMPORTANT_FIELDS.items() if not self.value(key)]
        return missing + ([] if self.nutrition else ["nutrition table"])


def field_key(label: str) -> str:
    """'Legal Name / Statement of Identity' -> 'legal_name_statement_of_identity'."""
    return re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")


_FIELD_FOR = {alias: key for key, aliases in FIELD_ALIASES.items() for alias in [key, *aliases]}
_NUTRIENT_FOR = {alias: name for name, aliases in NUTRIENT_ALIASES.items() for alias in aliases}


def canonical_field(label: str) -> str | None:
    """The checks' name for a sheet label, or None when it is not one they read.
    'Ingredients (English)' and 'Net Weight (g)' match too: a bracketed note is ignored."""
    for candidate in (label, re.sub(r"\(.*?\)", " ", label)):
        key = field_key(candidate.rstrip(":"))
        if key in _FIELD_FOR:
            return _FIELD_FOR[key]
    return None


def canonical_nutrient(label: str) -> str | None:
    return _NUTRIENT_FOR.get(field_key(re.sub(r"\(.*?\)", " ", label)))


# ---------------- the sheet as a grid of text cells ----------------


@dataclass
class Grid:
    """One worksheet as text: cells[(row, col)], both 1-based, empty cells left out."""

    title: str
    prefix: str  # "" for the first sheet, "Title!" for the others (cell references in findings)
    cells: dict[tuple[int, int], str] = field(default_factory=dict)

    @property
    def max_row(self) -> int:
        return max((r for r, _ in self.cells), default=0)

    @property
    def max_col(self) -> int:
        return max((c for _, c in self.cells), default=0)

    def text(self, row: int, col: int) -> str | None:
        return self.cells.get((row, col))

    def ref(self, row: int, col: int) -> str:
        return f"{self.prefix}{get_column_letter(col)}{row}"


def _clean(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    for dash in ("—", "–", "�"):  # sheets mix dash glyphs
        text = text.replace(dash, "-")
    return text or None


def _xlsx_grids(path: Path) -> list[Grid]:
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise SheetError(f"Could not open {path.name} as an Excel workbook: {exc}") from exc
    grids = []
    for i, ws in enumerate(wb.worksheets):
        grid = Grid(ws.title, "" if i == 0 else f"{ws.title}!")
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                # a cell formatted as a percentage stores 0.22 for "22%"
                if isinstance(value, (int, float)) and not isinstance(value, bool) and "%" in (cell.number_format or ""):
                    value = f"{value * 100:g}%"
                if (text := _clean(value)) is not None:
                    grid.cells[(cell.row, cell.column)] = text
        grids.append(grid)
    return grids


def _csv_grid(path: Path) -> Grid:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    # the delimiter that splits the most rows (values often contain commas: "Oats, Honey")
    delimiter = max(",;\t", key=lambda d: sum(len(row) > 1 for row in csv.reader(io.StringIO(text), delimiter=d)))
    grid = Grid(path.stem, "")
    for r, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        for c, value in enumerate(row, start=1):
            if (cleaned := _clean(value)) is not None:
                grid.cells[(r, c)] = cleaned
    return grid


# ---------------- fields ----------------


def _label_like(text: str) -> bool:
    """Short text with letters that is not itself a value (number, amount, sentence)."""
    return len(text) <= 60 and bool(re.search(r"[A-Za-z]", text)) and not AMOUNT.match(text)


def _table_rows(grid: Grid) -> set[int]:
    """Header rows of tables ("Product Name | Ingredients | Net Weight") and the value rows under them.
    A header row has two or more known labels and no value right next to any of them; a form row
    ("Country of Market | United States | ... | Product Name | Granola") has values beside its labels."""
    rows: set[int] = set()
    for row in range(1, grid.max_row + 1):
        known = [col for (r, col), text in grid.cells.items() if r == row and canonical_field(text)]
        beside = [grid.text(row, col + 1) for col in known]
        if len(known) >= 2 and all(t is None or canonical_field(t) for t in beside):
            rows |= {row, row + 1}
    return rows


def _label_columns(grid: Grid, rows: list[int]) -> list[int]:
    """Columns holding labels in a form layout: label-like text with a value right next to it, on at
    least two rows. Taken left to right; a label column claims the value and note columns after it."""
    columns, claimed = [], 0
    for col in range(1, grid.max_col + 1):
        if col <= claimed:
            continue
        hits = sum(1 for row in rows if (t := grid.text(row, col)) and _label_like(t) and grid.text(row, col + 1))
        if hits >= 2:
            columns.append(col)
            claimed = col + 2  # its value and note columns
    return columns


def _form_fields(grid: Grid, stop_row: int) -> dict[str, SheetField]:
    """Label / value / note triples in the sheet's label columns (table rows are read by _loose_fields)."""
    fields: dict[str, SheetField] = {}
    tables = _table_rows(grid)
    rows = [row for row in range(1, stop_row) if row not in tables]
    columns = _label_columns(grid, rows)
    for row in rows:
        for col in columns:
            label, value = grid.text(row, col), grid.text(row, col + 1)
            if not label or value is None:
                continue
            key = canonical_field(label) or field_key(label)
            if key == "field":  # a table header row ("Field | English (US) | Notes")
                continue
            fields[key] = SheetField(label=label, value=value, cell=grid.ref(row, col + 1), note=grid.text(row, col + 2))
    return fields


def _loose_fields(grid: Grid, found: dict[str, SheetField]) -> dict[str, SheetField]:
    """Fields the checks need that are not in a form layout: 'Label: value' in one cell, a label with
    its value a few cells to the right, or a header row with the values below it."""
    fields: dict[str, SheetField] = {}
    for (row, col), text in sorted(grid.cells.items()):
        if ":" in text and (head := text.split(":", 1)) and (key := canonical_field(head[0])) and head[1].strip():
            if key not in found and key not in fields:
                fields[key] = SheetField(label=head[0].strip(), value=head[1].strip(), cell=grid.ref(row, col))
            continue
        key = canonical_field(text) if _label_like(text) else None
        if not key or key in found or key in fields:
            continue
        right = next(((row, c) for c in range(col + 1, col + 4) if grid.text(row, c)), None)
        if right and not canonical_field(grid.text(*right)):
            fields[key] = SheetField(label=text, value=grid.text(*right), cell=grid.ref(*right))
        elif (below := grid.text(row + 1, col)) and not canonical_field(below):
            fields[key] = SheetField(label=text, value=below, cell=grid.ref(row + 1, col))
    return fields


# ---------------- nutrition ----------------


def _find_heading(grid: Grid) -> tuple[int, int] | None:
    return next(((r, c) for (r, c), t in sorted(grid.cells.items()) if t.lower() == NUTRITION_HEADING), None)


def _nutrient(name: str, amount: str, dv: str | None, cell: str) -> SheetNutrient:
    canonical = canonical_nutrient(name) or name
    if canonical == "Calories" and (kcal := re.search(r"(\d+(?:\.\d+)?)\s*kcal", amount, re.I)):
        amount = kcal.group(1)  # "1868kJ/445kcal" -> "445"
    return SheetNutrient(name=canonical, amount=amount, dv=None if dv in {None, "-"} else dv, cell=cell)


def _headed_nutrition(grid: Grid, row: int, col: int) -> tuple[str | None, list[SheetNutrient]]:
    """Rows below the heading: name, then the first basis amount, then its %DV."""
    basis = grid.text(row, col + 1)
    rows = []
    for r in range(row + 1, grid.max_row + 1):
        name, amount = grid.text(r, col), grid.text(r, col + 1)
        if not name or amount is None:
            break
        rows.append(_nutrient(name, amount, grid.text(r, col + 2), grid.ref(r, col + 1)))
    return basis, rows


def _loose_nutrition(grid: Grid) -> tuple[str | None, list[SheetNutrient]]:
    """No heading: every row that starts with a nutrient name followed by an amount (first one wins)."""
    rows: dict[str, SheetNutrient] = {}
    basis = None
    for (row, col), text in sorted(grid.cells.items()):
        name = canonical_nutrient(text)
        if not name or name in rows:
            continue
        amount_col = next((c for c in range(col + 1, col + 3) if (t := grid.text(row, c)) and AMOUNT.match(t)), None)
        if amount_col is None:
            continue
        amount = grid.text(row, amount_col)
        after = amount_col + 1
        if (unit := grid.text(row, after)) and UNIT.match(unit) and not amount[-1:].isalpha():
            amount, after = f"{amount}{unit}", after + 1  # amount and unit in separate columns
        dv = grid.text(row, after)
        rows[name] = _nutrient(text, amount, dv if dv and "%" in dv else None, grid.ref(row, amount_col))
        if basis is None and len(rows) == 1:
            basis = grid.text(row - 1, amount_col)  # the column heading above, e.g. "Per serving"
    return basis, list(rows.values())


# ---------------- reading a sheet ----------------


def read_content_sheet(path: str | Path) -> ApprovedSpec:
    path = Path(path)
    if not path.exists():
        raise SheetError(f"Content sheet not found: {path.name}")
    suffix = path.suffix.lower()
    if suffix == ".xls":
        raise SheetError(f"{path.name} is an old .xls workbook: save it as .xlsx (or .csv) and upload that.")
    grids = [_csv_grid(path)] if suffix == ".csv" else _xlsx_grids(path)

    fields: dict[str, SheetField] = {}
    basis, nutrition = None, []
    for grid in grids:
        heading = _find_heading(grid)
        stop = heading[0] if heading else grid.max_row + 1
        for key, value in _form_fields(grid, stop).items():
            fields.setdefault(key, value)  # the first sheet that has a field wins
        if not nutrition:
            basis, nutrition = _headed_nutrition(grid, *heading) if heading else _loose_nutrition(grid)
    for grid in grids:
        fields.update(_loose_fields(grid, fields))
    if not fields and not nutrition:
        raise SheetError(f"{path.name} has no labelled product fields or nutrition values that could be read")
    return ApprovedSpec(file_name=path.name, sheet_name=grids[0].title, fields=fields,
                        nutrition_basis=basis, nutrition=nutrition)
